#!/usr/bin/env python3
"""Mail merge: Excel rows + Word template (MERGEFIELDs) -> one PDF per row.

Subcommands:
  inspect  --input DIR                       -> JSON: files, columns, merge fields, mismatches
  run      --input DIR --work DIR [options]  -> PDFs in WORK, report in WORK/merge_report.json
"""
import argparse
import datetime as dt
import glob
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

from mailmerge import MailMerge
from openpyxl import load_workbook

LIBREOFFICE_PATHS = [
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    shutil.which("soffice") or "",
    shutil.which("libreoffice") or "",
]
WORD_APP = "/Applications/Microsoft Word.app"
CHROME_PATHS = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    shutil.which("google-chrome") or "",
    shutil.which("chromium") or "",
]


def fail(msg, **extra):
    print(json.dumps({"ok": False, "error": msg, **extra}, indent=2))
    sys.exit(1)


def find_inputs(input_dir):
    def pick(patterns, kind):
        files = sorted(
            f for p in patterns for f in glob.glob(os.path.join(input_dir, p))
            if not os.path.basename(f).startswith("~$")
        )
        if len(files) != 1:
            fail(f"expected exactly one {kind} in {input_dir}, found {len(files)}",
                 found=[os.path.basename(f) for f in files])
        return files[0]
    return pick(["*.xlsx", "*.xlsm"], "Excel file"), pick(["*.docx"], "Word template")


def format_value(value, date_format):
    if value is None:
        return ""
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime(date_format)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.2f}"
    return str(value).strip()


def read_rows(xlsx, sheet, date_format):
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    records = []
    for excel_row, values in enumerate(rows, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        rec = {h: format_value(v, date_format) for h, v in zip(header, values) if h}
        records.append((excel_row, rec))
    return ws.title, [h for h in header if h], records


def template_fields(docx):
    with MailMerge(docx) as doc:
        return sorted(doc.get_merge_fields())


def safe_filename(name):
    name = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", name).strip(" ._")
    return name or "document"


def inspect(args):
    xlsx, docx = find_inputs(args.input)
    sheet, columns, records = read_rows(xlsx, args.sheet, args.date_format)
    fields = template_fields(docx)
    print(json.dumps({
        "ok": True,
        "excel": os.path.basename(xlsx),
        "sheet": sheet,
        "template": os.path.basename(docx),
        "columns": columns,
        "merge_fields": fields,
        "fields_missing_in_excel": [f for f in fields if f not in columns],
        "columns_unused_by_template": [c for c in columns if c not in fields],
        "row_count": len(records),
        "sample_row": records[0][1] if records else None,
        "converter": detect_converter(),
    }, indent=2, ensure_ascii=False))


def detect_converter():
    for p in LIBREOFFICE_PATHS:
        if p and os.path.exists(p):
            return {"name": "libreoffice", "path": p}
    if os.path.exists(WORD_APP):
        try:
            import docx2pdf  # noqa: F401
            return {"name": "word", "path": WORD_APP}
        except ImportError:
            pass
    for p in CHROME_PATHS:
        if p and os.path.exists(p):
            return {"name": "chrome", "path": p}
    return None


def convert_libreoffice(soffice, docx_files, out_dir):
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, *docx_files],
        check=True, capture_output=True, timeout=600,
    )


def convert_word(docx_files, out_dir):
    from docx2pdf import convert
    for f in docx_files:
        convert(f, os.path.join(out_dir, os.path.splitext(os.path.basename(f))[0] + ".pdf"))


HTML_PAGE = """<!doctype html><html><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 2cm; }}
body {{ font-family: Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.4; }}
table {{ border-collapse: collapse; }} td, th {{ border: 1px solid #999; padding: 4px 8px; }}
img {{ max-width: 100%; }}
</style></head><body>{body}</body></html>"""


def convert_chrome(chrome, docx_files, out_dir):
    import mammoth
    for f in docx_files:
        base = os.path.splitext(os.path.basename(f))[0]
        with open(f, "rb") as fh:
            body = mammoth.convert_to_html(fh).value
        html_path = os.path.join(os.path.dirname(f), base + ".html")
        with open(html_path, "w", encoding="utf-8") as fh:
            fh.write(HTML_PAGE.format(body=body))
        subprocess.run(
            [chrome, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
             f"--print-to-pdf={os.path.join(out_dir, base + '.pdf')}",
             "file://" + os.path.abspath(html_path)],
            check=True, capture_output=True, timeout=120,
        )


def run(args):
    xlsx, docx = find_inputs(args.input)
    sheet, columns, records = read_rows(xlsx, args.sheet, args.date_format)
    fields = template_fields(docx)
    missing = [f for f in fields if f not in columns]
    if missing and not args.allow_missing:
        fail("template fields missing in Excel (use --allow-missing to leave them blank)",
             missing=missing)

    converter = detect_converter()
    if converter is None:
        fail("no PDF converter found: install LibreOffice (brew install --cask libreoffice)")

    os.makedirs(args.work, exist_ok=True)
    tmp = tempfile.mkdtemp(prefix="docx_", dir=args.work)
    planned, errors, used_names = [], [], set()

    for excel_row, rec in records:
        try:
            name = safe_filename(args.name_pattern.format(row=excel_row, **rec))
        except KeyError as e:
            errors.append({"row": excel_row, "error": f"name pattern column not found: {e}"})
            continue
        base, n = name, 2
        while base.lower() in used_names:
            base, n = f"{name}_{n}", n + 1
        used_names.add(base.lower())
        out_docx = os.path.join(tmp, base + ".docx")
        try:
            with MailMerge(docx) as doc:
                doc.merge(**{f: rec.get(f, "") for f in fields})
                doc.write(out_docx)
            planned.append({"row": excel_row, "file": base + ".pdf", "docx": out_docx, "data": rec})
        except Exception as e:  # keep going, report per row
            errors.append({"row": excel_row, "error": str(e)})

    docx_files = [p["docx"] for p in planned]
    if docx_files:
        try:
            if converter["name"] == "libreoffice":
                convert_libreoffice(converter["path"], docx_files, args.work)
            elif converter["name"] == "word":
                convert_word(docx_files, args.work)
            else:
                convert_chrome(converter["path"], docx_files, args.work)
        except subprocess.SubprocessError as e:
            fail(f"PDF conversion failed with {converter['name']}: {e}")

    generated = []
    for p in planned:
        if os.path.exists(os.path.join(args.work, p["file"])):
            generated.append({"row": p["row"], "file": p["file"], "data": p["data"]})
        else:
            errors.append({"row": p["row"], "error": "PDF was not produced by converter"})
    if not args.keep_docx:
        shutil.rmtree(tmp, ignore_errors=True)

    report = {
        "ok": not errors,
        "created": dt.datetime.now().isoformat(timespec="seconds"),
        "excel": os.path.basename(xlsx),
        "sheet": sheet,
        "template": os.path.basename(docx),
        "converter": converter["name"],
        "rows": len(records),
        "generated": generated,
        "errors": errors,
    }
    report_path = os.path.join(args.work, "merge_report.json")
    with open(report_path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=2, ensure_ascii=False)
    print(json.dumps({
        "ok": report["ok"], "report": report_path, "converter": converter["name"],
        "rows": len(records), "generated": len(generated), "errors": errors,
    }, indent=2, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "run"):
        p = sub.add_parser(name)
        p.add_argument("--input", default="input")
        p.add_argument("--sheet", default=None, help="sheet name (default: first sheet)")
        p.add_argument("--date-format", default="%d/%m/%Y")
        if name == "run":
            p.add_argument("--work", default="work")
            p.add_argument("--name-pattern", default="document_{row}",
                           help="file name from Excel columns, e.g. '{InvoiceNo}_{LastName}'")
            p.add_argument("--allow-missing", action="store_true")
            p.add_argument("--keep-docx", action="store_true")
    args = ap.parse_args()
    inspect(args) if args.cmd == "inspect" else run(args)


if __name__ == "__main__":
    main()
